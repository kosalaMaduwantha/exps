"""Excel workbook output."""

import pandas as pd
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter


def style_sheet(ws, money_cols=()):
    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    for row in ws.iter_rows():
        for c in row:
            c.font = font
    for c in ws[1]:
        c.font = bold
        c.alignment = Alignment(horizontal="center")
    for col_idx, col in enumerate(ws.columns, start=1):
        width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(width + 2, 10), 60)
    headers = [c.value for c in ws[1]]
    for name in money_cols:
        if name in headers:
            col = headers.index(name) + 1
            for r in range(2, ws.max_row + 1):
                ws.cell(row=r, column=col).number_format = "#,##0.00;(#,##0.00);-"
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def money_columns(statement_info):
    """Column names whose values are numeric (used for money formatting)."""
    return tuple(
        k for k, v in statement_info.items() if isinstance(v, (int, float))
    )


def write_workbook(out, transactions, statement_info):
    """Write the Transactions and Statement_Info sheets to an Excel file."""
    df_txn = pd.DataFrame(transactions)
    df_info = pd.DataFrame([statement_info], columns=list(statement_info.keys()))

    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        df_txn.to_excel(xw, sheet_name="Transactions", index=False)
        df_info.to_excel(xw, sheet_name="Statement_Info", index=False)
        style_sheet(xw.sheets["Transactions"], money_cols=("Amount",))
        style_sheet(
            xw.sheets["Statement_Info"],
            money_cols=money_columns(statement_info),
        )

    return df_txn, df_info
